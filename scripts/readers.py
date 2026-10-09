"""Independent-reader and determinism check.

    .venv/bin/python scripts/readers.py FILE...

For each file: rebuild it from its canonical JSON, open original and rebuild
with openpyxl (a separate Excel reader) and compare everything it exposes;
then convert twice and require identical artifacts.
"""

from __future__ import annotations

import sys
import tempfile
import warnings
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl  # noqa: E402

from provenance.canon import build_package, read_package  # noqa: E402
from provenance.canon import ConversionError  # noqa: E402
from provenance.cli import convert  # noqa: E402


def _value(v) -> str:
    """repr() of a cell value; formula objects get a stable form (their default repr has a memory address)."""
    if hasattr(v, "ref") and hasattr(v, "text"):  # ArrayFormula
        return f"ArrayFormula({v.ref!r}, {v.text!r})"
    if type(v).__name__ == "DataTableFormula":
        return "DataTableFormula(" + repr(sorted((k, str(x)) for k, x in vars(v).items())) + ")"
    return repr(v)


def snapshot(path: Path) -> dict:
    out = {}
    for data_only in (False, True):
        wb = openpyxl.load_workbook(path, data_only=data_only)
        for ws in wb.worksheets:
            cells = {}
            for row in ws.iter_rows():
                for c in row:
                    if c.value is not None or c.has_style:
                        cells[c.coordinate] = (_value(c.value), c.number_format, repr(c.font), repr(c.fill),
                                               repr(c.border), repr(c.alignment), repr(c.protection))
            out[(data_only, ws.title)] = {
                "cells": cells,
                "merged": sorted(str(r) for r in ws.merged_cells.ranges),
                "dims": ws.dimensions,
                "validations": [repr(v.sqref) + repr(v.formula1) for v in ws.data_validations.dataValidation],
                "cond": [str(r.sqref) for r in ws.conditional_formatting],
                "comments": sorted((c.coordinate, c.comment.text) for row in ws.iter_rows() for c in row if c.comment),
            }
        out[(data_only, "names")] = sorted((n, repr(d.attr_text)) for n, d in wb.defined_names.items())
        out[(data_only, "sheets")] = [(ws.title, ws.sheet_state) for ws in wb.worksheets]
    return out


def main() -> int:
    warnings.simplefilter("ignore")
    bad = 0
    for path in map(Path, sys.argv[1:]):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            problems = []
            if path.suffix.lower() in (".xlsx", ".xlsm"):
                rebuilt = tmp / ("rebuilt" + path.suffix)
                try:
                    rebuilt.write_bytes(build_package(read_package(path.read_bytes())))
                except ConversionError as exc:
                    print(f"SKIP {path.name[:60]:60} rejected: {str(exc)[:60]}")
                    continue
                try:
                    a = snapshot(path)
                except Exception as exc:
                    a, detail = None, f"(openpyxl cannot read the original either: {type(exc).__name__})"
                if a is not None:
                    try:
                        b = snapshot(rebuilt)
                        if a != b:
                            diffs = [k for k in a if a[k] != b.get(k)]
                            problems.append(f"openpyxl sees a difference in {diffs[:3]}")
                        detail = f"{sum(len(v['cells']) for k, v in a.items() if isinstance(v, dict) and 'cells' in v)} cell snapshots"
                    except Exception as exc:
                        problems.append(f"openpyxl reads the original but not the rebuild: {type(exc).__name__}: {exc}")
                        detail = ""
            else:
                detail = "(not a workbook: determinism only)"
            try:
                first = convert(path, tmp / "one", fuzz_cases=0)
                second = convert(path, tmp / "two", fuzz_cases=0)
            except ConversionError as exc:
                print(f"SKIP {path.name[:60]:60} rejected: {str(exc)[:60]}")
                continue
            for f in sorted((tmp / "one").iterdir()):
                g = tmp / "two" / f.name
                if f.name.endswith(".certificate.json"):
                    continue  # contains a timestamp; compared below without it
                if f.read_bytes() != g.read_bytes():
                    problems.append(f"{f.name} differs between two conversions")
            strip = lambda c: {k: v for k, v in c.items() if k not in ("created", "seal", "signature")}  # noqa: E731
            if strip(first) != strip(second):
                problems.append("certificate differs between two conversions")
        bad += bool(problems)
        print(f"{'OK  ' if not problems else 'FAIL'} {path.name[:60]:60} {detail}")
        for p in problems:
            print(f"     {p}")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
